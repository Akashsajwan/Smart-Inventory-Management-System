"""Tests for simulation visual charts and downloadable reports (Excel, PDF, CSV)."""
import io
import unittest
from openpyxl import load_workbook
import app as flask_app
from simulation_engine import run_simulation
from warehouse_demo import DEMO_PRODUCTS
from simulation_charts import (
    generate_all_simulation_charts,
    generate_simulation_pdf_report,
)


class SimulationChartsAndReportsTests(unittest.TestCase):
    def setUp(self):
        flask_app.app.config["TESTING"] = True
        self.client = flask_app.app.test_client()
        with flask_app.app.app_context():
            flask_app.db.create_all()

    def test_generate_all_six_simulation_charts(self):
        """Verify all 6 diverse simulation charts are generated as valid PNG buffers."""
        fixed = run_simulation([dict(p) for p in DEMO_PRODUCTS], duration_days=7, policy="fixed_threshold", random_seed=42)
        rop = run_simulation([dict(p) for p in DEMO_PRODUCTS], duration_days=7, policy="reorder_point", random_seed=42)

        charts = generate_all_simulation_charts(fixed, rop, start_day=1, end_day=7)
        expected_keys = [
            "cost_donut",
            "service_gauge",
            "inventory_trajectory",
            "daily_revenue",
            "sku_risk_scatter",
            "daily_fulfillment",
        ]
        for key in expected_keys:
            self.assertIn(key, charts)
            buf = charts[key]
            self.assertIsInstance(buf, io.BytesIO)
            # PNG signature bytes: 137 80 78 71 (b'\x89PNG')
            header = buf.getvalue()[:4]
            self.assertEqual(header, b"\x89PNG", f"Chart {key} should have valid PNG signature")
            self.assertGreater(len(buf.getvalue()), 1000, f"Chart {key} size should be substantial")

    def test_generate_simulation_pdf_report(self):
        """Verify executive PDF report is created with valid PDF header and substantial content."""
        fixed = run_simulation([dict(p) for p in DEMO_PRODUCTS], duration_days=7, policy="fixed_threshold", random_seed=42)
        rop = run_simulation([dict(p) for p in DEMO_PRODUCTS], duration_days=7, policy="reorder_point", random_seed=42)

        class MockRun:
            store_type = "electronics"
            random_seed = 42
            duration_days = 7

        pdf_buf = generate_simulation_pdf_report(
            MockRun(), MockRun(), fixed, rop, {}, {}, start_day=1, end_day=7
        )
        self.assertIsInstance(pdf_buf, io.BytesIO)
        pdf_bytes = pdf_buf.getvalue()
        # PDF magic bytes: %PDF-
        self.assertTrue(pdf_bytes.startswith(b"%PDF-"), "Should have valid PDF header")
        self.assertGreater(len(pdf_bytes), 10000, "PDF document should contain multiple formatted pages and charts")

    def test_simulation_run_and_download_routes(self):
        """Verify running simulation via POST and downloading reports in xlsx, pdf, and csv formats."""
        # Run simulation via POST
        post_resp = self.client.post("/simulation", data={
            "store_type": "electronics",
            "duration_days": "7",
            "scenario": "normal",
            "random_seed": "42",
        }, follow_redirects=True)
        self.assertEqual(post_resp.status_code, 200)
        html = post_resp.get_data(as_text=True)
        self.assertIn("Simulation Visual Analytics", html)
        self.assertIn("Download Excel Report (.xlsx)", html)
        self.assertIn("Download Visual PDF Report (.pdf)", html)
        self.assertIn("Download CSV Summary (.csv)", html)
        self.assertIn("cost-donut-chart", html)
        self.assertIn("service-gauge-chart", html)
        self.assertIn("inventory-trajectory-chart", html)
        self.assertIn("revenue-comparison-chart", html)
        self.assertIn("sku-risk-scatter-chart", html)
        self.assertIn("sales-comparison-chart", html)
        # Check explanatory lines and callouts in Web UI
        self.assertIn("Chart Explanation & Guide Lines", html)
        self.assertIn("95% Target Line", html)
        self.assertIn("Safety Stock Buffer Line", html)
        self.assertIn("Avg Revenue Baseline", html)
        self.assertIn("Crosshair Lines", html)
        self.assertIn("Avg Demand Target", html)

        # Check ML future sales forecasting UI section and chart
        self.assertIn("ml-forecast-section", html)
        self.assertIn("ml-forecast-chart", html)
        self.assertIn("Machine Learning Future Sales", html)
        self.assertIn("Forward SKU Stockout Depletion Risk Matrix", html)

        # Get latest run IDs from DB
        with flask_app.app.app_context():
            runs = flask_app.SimulationRun.query.order_by(flask_app.SimulationRun.id.desc()).limit(2).all()
            self.assertEqual(len(runs), 2)
            # Find fixed and rop run ids
            fixed_run = next(r for r in runs if r.policy_name == "fixed_threshold")
            rop_run = next(r for r in runs if r.policy_name == "reorder_point")

        # 1. Test Excel download with Visual Dashboard
        xlsx_resp = self.client.get(f"/simulation/report/{fixed_run.id}/{rop_run.id}/xlsx?start_day=1&end_day=7")
        self.assertEqual(xlsx_resp.status_code, 200)
        self.assertEqual(xlsx_resp.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("smart_warehouse_report_days_1_7.xlsx", xlsx_resp.headers["Content-Disposition"])

        wb = load_workbook(io.BytesIO(xlsx_resp.data))
        self.assertIn("Visual Dashboard", wb.sheetnames)
        self.assertIn("Management Summary", wb.sheetnames)
        self.assertIn("Daily Metrics", wb.sheetnames)
        self.assertIn("Product Stock & Restock", wb.sheetnames)
        self.assertIn("ML Sales Forecast", wb.sheetnames)

        dash_ws = wb["Visual Dashboard"]
        self.assertGreaterEqual(len(dash_ws._images), 7, "Visual Dashboard should embed all 7 chart images including ML forecast")
        # Verify explanatory text lines below each chart in Visual Dashboard
        self.assertIn("Donut rings display cost proportion", str(dash_ws["B22"].value))
        self.assertIn("95% Target Line", str(dash_ws["M22"].value))
        self.assertIn("Safety Stock Buffer Line", str(dash_ws["B42"].value))
        self.assertIn("Average Daily Revenue Baseline", str(dash_ws["M42"].value))
        self.assertIn("Crosshairs", str(dash_ws["B62"].value))
        self.assertIn("Average Demand Target Line", str(dash_ws["M62"].value))
        self.assertIn("Predictive Ridge Regression", str(dash_ws["B83"].value))

        # Check ML Sales Forecast sheet contents
        ml_ws = wb["ML Sales Forecast"]
        self.assertIn("MACHINE LEARNING 7-DAY FORWARD SALES", str(ml_ws["A1"].value))
        self.assertIn("Ridge Regression", str(ml_ws["B2"].value))

        # 2. Test PDF download with Executive Graphical Report
        pdf_resp = self.client.get(f"/simulation/report/{fixed_run.id}/{rop_run.id}/pdf?start_day=1&end_day=7")
        self.assertEqual(pdf_resp.status_code, 200)
        self.assertEqual(pdf_resp.mimetype, "application/pdf")
        self.assertIn("smart_warehouse_visual_report_days_1_7.pdf", pdf_resp.headers["Content-Disposition"])
        self.assertTrue(pdf_resp.data.startswith(b"%PDF-"), "Downloaded PDF must start with %PDF- header")
        self.assertGreater(len(pdf_resp.data), 20000)

        # 3. Test CSV download
        csv_resp = self.client.get(f"/simulation/report/{fixed_run.id}/{rop_run.id}/csv?start_day=1&end_day=7")
        self.assertEqual(csv_resp.status_code, 200)
        self.assertEqual(csv_resp.mimetype, "text/csv")
        self.assertIn("smart_warehouse_report_days_1_7.csv", csv_resp.headers["Content-Disposition"])
        csv_text = csv_resp.data.decode("utf-8-sig")
        self.assertIn("SMART WAREHOUSE SIMULATION REPORT", csv_text)
        self.assertIn("MACHINE LEARNING 7-DAY FORWARD SALES FORECAST (DAYS 8-14)", csv_text)
        self.assertIn("Product-Level Forward Stockout Risk Assessment", csv_text)

    def test_stress_scenario_charts_and_reports(self):
        """Verify chart generation works seamlessly during high-stress scenarios with lost sales."""
        products = [dict(p) for p in DEMO_PRODUCTS]
        for item in products:
            item["avg_daily_demand"] = max(1, round(item["avg_daily_demand"] * 1.8))
            item["demand_stddev"] = max(1, round(item["demand_stddev"] * 1.5))
            item["lead_time_days"] += 2
            item["initial_stock"] = max(1, item["initial_stock"] // 2)

        fixed = run_simulation(products, duration_days=7, policy="fixed_threshold", random_seed=42)
        rop = run_simulation(products, duration_days=7, policy="reorder_point", random_seed=42)

        # Verify lost units occurred in stress test
        self.assertGreater(fixed["summary"]["total_lost_units"], 0)

        charts = generate_all_simulation_charts(fixed, rop, start_day=1, end_day=7)
        self.assertIn(len(charts), (6, 7))
        for key, buf in charts.items():
            self.assertTrue(buf.getvalue().startswith(b"\x89PNG"))
